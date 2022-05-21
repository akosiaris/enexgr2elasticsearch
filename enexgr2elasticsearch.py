'''
Populate elasticsearch with Greek Energy Exchange Group free data

Copyright Alexandros Kosiaris 2022
'''

import argparse
import io
import json
import logging
import os
import sys
import warnings
from datetime import datetime, timedelta
from logging.handlers import HTTPHandler
from urllib.parse import urljoin

import ecs_logging
import requests
import structlog
from openpyxl import load_workbook as _load_workbook
from pytz import timezone
from requests.auth import HTTPBasicAuth

VERSION = '0.1'
ELECTRICITY_MARKETS_META_DATA = {
    'RESULTS': {
        'index': 'enexgr_electricity_market_results',
        'base_urls': {
            'DAM':   '20126/200106/%s_EL-DAM_Results_EN_v01.xlsx',
            'LIDA1': '20126/235155/%s_EL-LIDA1_Results_EN_v01.xlsx',
            'LIDA2': '20126/263261/%s_EL-LIDA2_Results_EN_v01.xlsx',
            'LIDA3': '20126/263280/%s_EL-LIDA3_Results_EN_v01.xlsx',
            'CRIDA1': '20126/853663/%s_EL-CRIDA1_Results_EN_v01.xlsx',
            'CRIDA2': '20126/853680/%s_EL-CRIDA2_Results_EN_v01.xlsx',
            'CRIDA3': '20126/853704/%s_EL-CRIDA3_Results_EN_v01.xlsx',
        },
    },
    'CURVES': {
        'index': 'enexgr_electricity_market_curves',
        'base_urls': {
            'DAM': '20126/200034/%s_EL-DAM_AggrCurves_EN_v01.xlsx',
            'CRIDA1': '20126/853660/%s_EL-CRIDA1_AggrCurves_EN_v01.xlsx',
            'CRIDA2': '20126/853695/%s_EL-CRIDA2_AggrCurves_EN_v01.xlsx',
            'CRIDA3': '20126/853701/%s_EL-CRIDA3_AggrCurves_EN_v01.xlsx',
        },
    },
    'BLOCK_ORDERS': {
        'index': 'enexgr_electricity_market_block_orders',
        'base_urls': {
            'DAM': '20126/270103/%s_EL-DAM_BLKORDRs_EN_v01.xlsx',
        },
    },
    'HOURLY_DAILY_MCPS': {
        # The hourly_daily_mcp index is a derivative one, so we only have a name
        'index': 'enexgr_electricity_market_hourly_daily_mcps',
    },
}

GAS_MARKETS_META_DATA = {
    'NGAS_Results': {
        'index': 'enexgr_gas_market_results',
        'base_urls': {
            'NGAS_DOL': '20126/997118/%s_NGAS_DOL_EN_v01.xlsx',
        },
    },
    'Auctions_Details': {
        'index': 'enexgr_gas_market_details',
    },
    'Announcements': {
        'index': 'enexgr_gas_market_announcements',
    },
}

BULK_ENDPOINT = '/_bulk'
TZ = timezone('Europe/Athens')
DELTA = timedelta(days=1)

BASE_ENEX_URL = 'https://www.enexgroup.gr/documents'

def put_to_elastic(data: str, url: str, elastic_info: dict) -> bool:
    '''
    PUT to elasticsearch
    '''
    user = elastic_info.get('user')
    password = elastic_info.get('password')
    if user and password:
        auth=HTTPBasicAuth(user, password)
    else:
        auth=None

    logger.debug('Data to be PUT to elasticsearch', data=data)
    try:
        response = requests.put(
            url,
            headers={
                'Content-Type': 'application/json',
            },
            data=data.encode('utf-8'),
            auth=auth)
        if response.status_code != 200:
            logger.error('Elasticsearch error response',
                    status_code=response.status_code,
                    body=response.content.decode())
            return False
    except requests.exceptions.ConnectionError as exc:
        logger.error('Connection failed', exc_info=exc)

    logger.debug('Data PUT successfully to elasticsearch', size=len(data))
    return True


def post_to_bulk_elastic(data: str, url: str, elastic_info: dict) -> bool:
    '''
    Post to elasticsearch
    '''
    user = elastic_info.get('user')
    password = elastic_info.get('password')
    if user and password:
        auth=HTTPBasicAuth(user, password)
    else:
        auth=None

    logger.debug('Data to be POSTed to elasticsearch', data=data)
    try:
        response = requests.post(
            url,
            headers={
                'Content-Type': 'application/x-ndjson',
            },
            data=data.encode('utf-8'),
            auth=auth)
        if response.status_code != 200:
            logger.error('Elasticsearch error response',
                    status_code=response.status_code,
                    body=response.content.decode())
            return False
    except requests.exceptions.ConnectionError as exc:
        logger.error('Connection failed', exc_info=exc)

    logger.debug('Bulk data indexed succesfully', size=len(data))
    return True


def fetch_new_xlsx(url: str) -> str:
    '''
    Fetch and validate a new xlsx from enexgroup.gr
    '''

    resp = requests.get(url)
    if resp.status_code == 200:
        return resp.content
    if resp.status_code == 404 or resp.status_code == 403:
        return None
    raise RuntimeError('Failed to fetch: %s' % resp.status_code)


def convert_electricity_market_results_workbook(xlsx: str) -> tuple:
    '''
    Converts the data from an enexgroup market result xlsx file to a ready for
    elasticsearch bulk API POST string
    '''

    workbook = load_workbook(xlsx)
    rows = workbook.active.rows

    tmp = next(rows)
    header = [x.value for x in tmp]

    ret = ''
    hourly_mcps = set()
    for row in rows:
        tmp = [x.value for x in row]
        data = dict(zip(header, tmp))
        delivery_timestamp = datetime.fromisoformat(data['DELIVERY_MTU'])
        data['DELIVERY_MTU'] = TZ.localize(delivery_timestamp).isoformat()
        pub_timestamp = datetime.fromisoformat(data['PUB_TIME'])
        data['PUB_TIME'] = TZ.localize(pub_timestamp).isoformat()
        ret += '{ "index": { "_index": "%s", "_id": "%s-%s-%s-%s-%s-%s-%s" } }' % (
                ELECTRICITY_MARKETS_META_DATA['RESULTS']['index'],
                data['TARGET'],
                data['BIDDING_ZONE_DESCR'],
                data['SIDE_DESCR'],
                data['DDAY'],
                data['ASSET_DESCR'],
                data['CLASSIFICATION'],
                data['DELIVERY_MTU'])
        ret += '\n' + json.dumps(data) + '\n'
        if data['TARGET'] == 'DAM':
            hourly_mcps.add((data['MCP'], data['DELIVERY_MTU']))
    return (ret, hourly_mcps)


def calculate_electricity_hourly_daily_mcps(hourly_mcps: set) -> str:
    '''
    Calculate hourly/daily MCPs and send to dedicated index
    '''
    ret = ''
    if len(hourly_mcps) > 0:
        for mcp in hourly_mcps:
            ret += '{ "index": { "_index": "%s", "_id": "DAM-%s-MCP-HOURLY" } }' % (
                    ELECTRICITY_MARKETS_META_DATA['HOURLY_DAILY_MCPS']['index'],
                    mcp[1])
            ret += '\n{ "HOURLY_MCP": %s, "DELIVERY_MTU": "%s" }\n' % mcp
        daily_mcp = sum([x[0] for x in hourly_mcps])/len(hourly_mcps)
        dtime = min([x[1] for x in hourly_mcps])
        ret += '{ "index": { "_index": "%s", "_id": "DAM-%s-MCP-DAILY" } }' % (
                ELECTRICITY_MARKETS_META_DATA['HOURLY_DAILY_MCPS']['index'],
                dtime)
        ret += '\n{ "DAILY_MCP": %s, "DELIVERY_MTU": "%s" }\n' % (daily_mcp, dtime)
    return ret


def convert_electricity_curves_workbook(xlsx: str) -> str:
    '''
    Convert the data from an enexgroup aggregated curves result xlsx file to a
    ready for elasticsearch bulk API POST string
    '''

    workbook = load_workbook(xlsx)
    rows = workbook.active.rows

    tmp = next(rows)
    header = [x.value for x in tmp]

    ret = ''
    for row in rows:
        tmp = [x.value for x in row]
        data = dict(zip(header, tmp))
        # NOTE: Yes, the date format is inconsistent across time
        data['DELIVERY_MTU'] = data['DELIVERY_MTU'].replace('/', '-')
        data['PUB_TIME'] = data['PUB_TIME'].replace('/', '-')
        delivery_timestamp = datetime.fromisoformat(data['DELIVERY_MTU'])
        data['DELIVERY_MTU'] = TZ.localize(delivery_timestamp).isoformat()
        pub_timestamp = datetime.fromisoformat(data['PUB_TIME'])
        data['PUB_TIME'] = TZ.localize(pub_timestamp).isoformat()
        ret += '{ "index": { "_index": "%s", "_id": "%s-%s-%s-%s-%s" } }' % (
                ELECTRICITY_MARKETS_META_DATA['CURVES']['index'],
                data['TARGET'],
                data['SIDE_DESCR'],
                data['DDAY'],
                data['AA'],
                data['DELIVERY_MTU'])
        ret += '\n' + json.dumps(data) + '\n'
    return ret


def convert_electricity_blockorders_workbook(xlsx: str) -> str:
    '''
    Convert the data from an enexgroup block orders result xlsx file to a
    ready for elasticsearch bulk API POST string
    '''

    workbook = load_workbook(xlsx)
    rows = workbook.active.rows

    tmp = next(rows)
    header = [x.value for x in tmp]

    ret = ''
    for row in rows:
        tmp = [x.value for x in row]
        data = dict(zip(header, tmp))
        delivery_timestamp = datetime.fromisoformat(data['DELIVERY_MTU'])
        data['DELIVERY_MTU'] = TZ.localize(delivery_timestamp).isoformat()
        pub_timestamp = datetime.fromisoformat(data['PUB_TIME'])
        data['PUB_TIME'] = TZ.localize(pub_timestamp).isoformat()
        ret += '{ "index": { "_index": "%s", "_id": "%s-%s-%s-%s-%s-%s" } }' % (
                ELECTRICITY_MARKETS_META_DATA['BLOCK_ORDERS']['index'],
                data['TARGET'],
                data['BIDDING_ZONE_DESCR'],
                data['SIDE_DESCR'],
                data['DDAY'],
                data['CLASSIFICATION'],
                data['DELIVERY_MTU'])
        ret += '\n' + json.dumps(data) + '\n'
    return ret


def convert_gas_workbook(xlsx: str) -> str:
    '''
    Convert the data from an enexgroup NGAS DOL result xlsx file to a
    ready for elasticsearch bulk API POST string
    '''

    workbook = load_workbook(xlsx)

    ret = ''
    for worksheet in workbook.worksheets:
        rows = worksheet.rows
        tmp = next(rows)
        # The if is there cause of empty columns with no name
        header = [x.value.strip() for x in tmp if x.value]

        for row in rows:
            tmp = [x.value for x in row if x.value]
            data = dict(zip(header, tmp))
            data['Trading Date'] = TZ.localize(data['Trading Date']).isoformat()
            ret += '{ "index": { "_index": "%s", "_id": "%s-%s-%s" } }' % (
                    GAS_MARKETS_META_DATA[worksheet.title]['index'],
                    data['Trading Date'],
                    data['Trading Series'],
                    data['Contract'])
            ret += '\n' + json.dumps(data) + '\n'
    return ret


def create_elasticsearch_indices(elastic_admin_info: dict, shards: int,
        replicas: int) -> bool:
    '''
    Create the elasticsearch indices alongside mappings
    '''

    electricity_indices = map(lambda x: x[1]['index'], ELECTRICITY_MARKETS_META_DATA.items())
    gas_indices = map(lambda x: x[1]['index'], GAS_MARKETS_META_DATA.items())
    settings = {
        "settings": {
            "number_of_shards": shards,
            "number_of_replicas": replicas,
        }
    }
    indices = list(electricity_indices) + list(gas_indices)
    for idx in indices:
        with open('%s.index' % idx, 'r') as fil:
            # Load the mappings
            data = json.load(fil)
            data.update(settings)
            url = urljoin(elastic_admin_info['host'], idx)
            if put_to_elastic(json.dumps(data), url, elastic_admin_info):
                logger.info('index created succesfully', index=idx)
            else:
                logger.warning('index creation failed', index=idx)
                return False
    return True


def get_xlsx(cache: str, filepath: str) -> io.BytesIO:
    '''
    Get xlsx from cache or fetch from internet
    '''
    url = BASE_ENEX_URL + '/' + filepath
    _, filename = os.path.split(filepath)
    cache_path = None
    if cache:
        cache_path = os.path.join(cache, filename)
        try:
            tmp = open(cache_path, 'rb').read()
            logger.info('Cache-hit', xlsx=filename)
            return io.BytesIO(tmp)
        except FileNotFoundError:
            logger.info('Cache-miss', xlsx=filename)
    tmp = fetch_new_xlsx(url)
    if tmp:
        logger.info('Successful download', xlsx=filename)
        # Write to cache if enabled
        if cache_path:
            with open(cache_path, 'wb') as cache_file:
                cache_file.write(tmp)
        return io.BytesIO(tmp)
    return None


def load_workbook(xlsx: str):
    '''
    Overriding load_workbook to catch exceptions and silence warnings
    '''

    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            workbook = _load_workbook(
                    filename=xlsx,
                    read_only=False)
    except Exception as exc:
        logger.error('Loading xlsx failed', xlsx=xlsx)
        raise exc
    return workbook


def setup_logging(args):
    '''
    Setting up logging function
    '''

    if args.verbose == 1:
        level = logging.INFO
    elif args.verbose > 1:
        level = logging.DEBUG
    else:
        level = logging.WARN
    processors = [
            # If log level is too low, abort pipeline and throw away log entry.
            structlog.stdlib.filter_by_level,
            # Add the name of the logger to event dict.
            structlog.stdlib.add_logger_name,
            # If the "stack_info" key in the event dict is true, remove it and
            # render the current stack trace in the "stack" key.
            structlog.processors.StackInfoRenderer(),
            # If some value is in bytes, decode it to a unicode str.
            structlog.processors.UnicodeDecoder(),
            # Add callsite parameters.
            structlog.processors.CallsiteParameterAdder(
                parameters=[
                    structlog.processors.CallsiteParameter.FUNC_NAME,
                    structlog.processors.CallsiteParameter.LINENO
                ]
            ),
    ]
    if hasattr(sys.stdout, 'isatty') and sys.stdout.isatty():
        # Running in a terminal, assume dev and setup nice stuff
        processors += [
            structlog.stdlib.add_log_level,
            structlog.processors.TimeStamper(),
            structlog.dev.set_exc_info,
            structlog.dev.ConsoleRenderer()
        ]
        # And we want a logging level of INFO anyway
        level = logging.INFO
    else:
        # Running under some supervisor, be more production-y
        processors += [
            # If the "exc_info" key in the event dict is either true or a
            # sys.exc_info() tuple, remove "exc_info" and render the exception
            # with traceback into the "exception" key.
            structlog.processors.format_exc_info,
        ]
        handler = logging.StreamHandler(sys.stdout)
        if args.ecs_logging:
            processors += [
                ecs_logging.StructlogFormatter()
            ]
            handler = ElasticSearchLogHandler(
                host='localhost:9200',
                url='/logs-enexgr2elasticsearch-1/_doc',
                method='POST',
                secure=False,
                credentials=(args.user, args.password),
            )
        else:
            processors += [
                structlog.stdlib.add_log_level,
                # Add a timestamp in ISO 8601 format.
                structlog.processors.TimeStamper(fmt="iso"),
                structlog.processors.KeyValueRenderer()
            ]


    logging.basicConfig(
        format="%(message)s",
        handlers=[handler],
        level=level,
    )
    structlog.configure(
        processors=processors,
        wrapper_class=structlog.stdlib.BoundLogger,
        logger_factory=structlog.stdlib.LoggerFactory(),
        cache_logger_on_first_use=True,
    )
    return structlog.get_logger(__name__)


class ElasticSearchLogHandler(HTTPHandler):
    '''
    Class logging to elasticsearch
    '''

    def emit(self, record):
        """
        Emit a record.
        Send the record to the web server as a application/json
        """
        try:
            msg = self.format(record)
            host = self.host
            h = self.getConnection(host, self.secure)
            url = self.url
            if self.method == "GET":
                raise ValueError("No GET please")
            h.putrequest(self.method, url)
            # support multiple hosts on one IP address...
            # need to strip optional :port from host, if present
            i = host.find(":")
            if i >= 0:
                host = host[:i]
            # See issue #30904: putrequest call above already adds this header
            # on Python 3.x.
            # h.putheader("Host", host)
            if self.method == "POST":
                h.putheader("Content-type",
                            "application/json")
                h.putheader("Content-length", str(len(msg)))
            if self.credentials:
                import base64
                s = ('%s:%s' % self.credentials).encode('utf-8')
                s = 'Basic ' + base64.b64encode(s).strip().decode('ascii')
                h.putheader('Authorization', s)
            h.endheaders()
            if self.method == "POST":
                h.send(msg.encode('utf-8'))
            h.getresponse()    #can't do anything with the result
        except Exception:
            self.handleError(record)


def main():
    '''
    Main function
    '''
    global logger

    parser = argparse.ArgumentParser(
            prog='enexgr.py',
            description='Fetch enexgroup.gr data and put into Elasticsearch')
    parser.add_argument('-H',
                        '--host',
                        dest='host',
                        default='http://localhost:9200',
                        help='URL pointing to the elasticsearch cluster')
    parser.add_argument('-u',
                        '--user',
                        dest='user',
                        help='Elasticsearch user to write data')
    parser.add_argument('-p',
                        '--password',
                        dest='password',
                        help='Elasticsearch user password')
    parser.add_argument('--admin-user',
                        dest='admin_user',
                        help='Elasticsearch admin user to create the indices')
    parser.add_argument('--admin-password',
                        dest='admin_password',
                        help='Elasticsearch admin user password')
    parser.add_argument('--create-indices',
                        action='store_true',
                        dest='create_indices',
                        help='Create Elasticsearch indices. Requires elasticsearch admin access')
    parser.add_argument('--create-indices-shards',
                        default=1,
                        dest='shards',
                        help='''Number of shards for each newly created
                        Elasticsearch index. Used only when creating indices.
                        Defaults to 1. Requires elasticsearch admin access''')
    parser.add_argument('--create-indices-replicas',
                        default=0,
                        dest='replicas',
                        help='''Number of replicas for each newly created
                        Elasticsearch index. Used only when creating indices.
                        Defaults to 0. Requires elasticsearch admin access''')
    parser.add_argument('-s',
                        '--start',
                        dest='start',
                        default=(datetime.now().date()-DELTA).strftime('%Y-%m-%d'),
                        help='The start date. YYYY-MM-DD format. Defaults to yesterday')
    parser.add_argument('-e',
                        '--end',
                        dest='end',
                        default=(datetime.now().date()+DELTA).strftime('%Y-%m-%d'),
                        help='The end date. YYYY-MM-DD format. Defaults to tomorrow')
    parser.add_argument('-c',
                        '--cache',
                        dest='cache',
                        help='Directory to be used as R/W cache for .xlsx files')
    parser.add_argument('-v',
                        '--verbose',
                        action='count',
                        default=0,
                        dest='verbose',
                        help='Increase verbosity. May be specified multiple times')
    parser.add_argument('--ecs-logging',
                        action='store_true',
                        default=False,
                        dest='ecs_logging',
                        help='Use Elastic Common Schema logging')
    parser.add_argument('--version',
                        action='version',
                        version='%(prog)s ' + VERSION)
    args = parser.parse_args()

    # Let's setup logging first
    logger = setup_logging(args)

    elastic_info = dict(
            host = args.host,
            user = args.user,
            password = args.password)
    elastic_admin_info = dict(
            host = args.host,
            user = args.admin_user,
            password = args.admin_password)

    if args.create_indices:
        if not create_elasticsearch_indices(
            elastic_admin_info,
            args.shards,
            args.replicas):
            logger.critical('Failed to create indices despite being asked to',
                    shards=args.shards,
                    replicas=args.replicas)
            return 1

    start_date = datetime.strptime(args.start, '%Y-%m-%d')
    end_date = datetime.strptime(args.end, '%Y-%m-%d')

    day_count = (end_date - start_date).days + 1

    bulk_url = urljoin(elastic_info['host'], BULK_ENDPOINT)
    for delta in range(0, day_count):
        date = start_date + timedelta(days=delta)
        markets_meta_data = {}
        markets_meta_data.update(ELECTRICITY_MARKETS_META_DATA)
        markets_meta_data.update(GAS_MARKETS_META_DATA)
        for category, data in markets_meta_data.items():
            try:
                base_urls = data['base_urls']
            except KeyError:
                logger.debug('Category missing base_urls, skipping',
                        category=category)
                continue
            for market, base_url in base_urls.items():
                # In 2021-09-22 LIDAs were renamed to CRIDAs. Don't try to fetch
                # LIDAs after this time and CRIDAs before this time
                if date > datetime(2021, 9, 21) and market.startswith('LIDA'):
                    continue
                if date <= datetime(2021, 9, 21) and market.startswith('CRIDA'):
                    continue

                filepath = base_url % date.strftime('%Y%m%d')
                xlsx = get_xlsx(args.cache, filepath)
                if xlsx:
                    logger.info('Successful xlsx fetch',
                        xlsx_date=date.isoformat(), category=category, market=market)
                    if category == 'RESULTS':
                        data, hourly_mcps = convert_electricity_market_results_workbook(xlsx)
                        data = data + calculate_electricity_hourly_daily_mcps(hourly_mcps)
                    if category == 'CURVES':
                        data = convert_electricity_curves_workbook(xlsx)
                    if category == 'BLOCK_ORDERS':
                        data = convert_electricity_blockorders_workbook(xlsx)
                    if category == 'NGAS_Results':
                        data = convert_gas_workbook(xlsx)
                    logger.info('xlsx to json done',
                            xlsx_date=date.isoformat(), category=category, market=market)
                    if post_to_bulk_elastic(data, bulk_url, elastic_info):
                        logger.info('Posted to bulk API',
                            xlsx_date=date.isoformat(), category=category, market=market)


if __name__ == '__main__':
    logger = None
    main()
